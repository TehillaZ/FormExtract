"""The records list (filters, totals, paging) and the export: one set of filters, so they must always agree."""
import csv
import io
import itertools

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import app.db.session as session_mod
from app.api.main import app
from app.db import repo
from app.pipeline.extract import build_record
from app.schemas import RawExtraction, RawValue as V

counter = itertools.count(1)


@pytest.fixture()
def client():
    global counter
    counter = itertools.count(1)        # each test numbers its records from 1, so numbers like 930001 are predictable
    session_mod.init_db("sqlite://")
    app.state.providers = None
    with TestClient(app) as c:
        yield c


def seed(first="דנה", last="כהן", sender="כללית", received="18/06/2026", id_="011223344", status="pending") -> str:
    """Save one record with its own numbers (so approving never collides) and return its id."""
    n = next(counter)
    raw = RawExtraction(
        first_name=V(value=first, source="printed"), last_name=V(value=last, source="printed"),
        id_numbers=[V(value=id_, source="printed")],
        central_numbers=[V(value=f"93{n:04d}", source="sticker")],
        commitment_numbers=[V(value=f"7305{n:05d}", source="printed")],
        sender=V(value=sender, source="printed", highlighted=True), stamp_date=V(value=received, source="stamp"))
    with session_mod._factory() as s:
        rec = repo.save_record(s, build_record(raw, []))
        if status == "approved":
            repo.approve_record(s, rec.id)
        return str(rec.id)


def total(response) -> int:
    return int(response.headers["X-Total-Count"])


class TestListAndTotals:
    def test_the_total_counts_every_match_not_just_the_page(self, client):
        for _ in range(5):
            seed()
        page = client.get("/records", params={"limit": 2})
        assert len(page.json()) == 2 and total(page) == 5

    def test_the_last_page_is_short_and_the_total_is_unchanged(self, client):
        for _ in range(5):
            seed()
        page = client.get("/records", params={"limit": 2, "offset": 4})
        assert len(page.json()) == 1 and total(page) == 5

    def test_pages_do_not_overlap_and_together_cover_everything(self, client):
        for _ in range(5):
            seed()
        ids = [r["id"] for off in (0, 2, 4) for r in client.get("/records", params={"limit": 2, "offset": off}).json()]
        assert len(ids) == 5 == len(set(ids))

    def test_the_total_follows_the_filters(self, client):
        seed(status="approved")
        seed()
        seed()
        assert total(client.get("/records", params={"status": "approved"})) == 1
        assert total(client.get("/records", params={"status": "pending"})) == 2
        assert total(client.get("/records")) == 3

    def test_an_empty_result_has_a_zero_total(self, client):
        response = client.get("/records", params={"q": "nobody"})
        assert response.json() == [] and total(response) == 0

    def test_a_negative_offset_is_rejected(self, client):
        assert client.get("/records", params={"offset": -1}).status_code == 422

    def test_a_bad_status_is_rejected(self, client):
        assert client.get("/records", params={"status": "weird"}).status_code == 422


class TestFilters:
    def test_search_finds_by_name_id_central_and_commitment_number(self, client):
        seed(first="רונן", last="פלג", id_="022334455")
        seed(first="דנה", last="כהן")
        names = lambda q: [r["last_name"]["value"] for r in client.get("/records", params={"q": q}).json()]  # noqa: E731
        assert names("פלג") == ["פלג"]                 # last name
        assert names("רונ") == ["פלג"]                 # part of a first name
        assert names("022334455") == ["פלג"]           # ID
        assert names("930001") == ["פלג"]              # central number of the first record
        assert names("730500002") == ["כהן"]           # commitment number of the second
        assert len(names("")) == 2                     # empty search = everything

    def test_filter_by_sender(self, client):
        seed(sender="כללית")
        seed(sender="מכבי")
        assert [r["sender"]["value"] for r in client.get("/records", params={"sender": "מכבי"}).json()] == ["מכבי"]

    def test_filter_by_received_date_range_includes_both_ends(self, client):
        for day in ("10/06/2026", "15/06/2026", "20/06/2026"):
            seed(received=day)
        got = lambda **p: sorted(r["received_date"]["value"] for r in client.get("/records", params=p).json())  # noqa: E731
        assert got(**{"from": "2026-06-15"}) == ["2026-06-15", "2026-06-20"]
        assert got(to="2026-06-15") == ["2026-06-10", "2026-06-15"]
        assert got(**{"from": "2026-06-15", "to": "2026-06-15"}) == ["2026-06-15"]

    def test_filters_combine(self, client):
        seed(sender="כללית", status="approved")
        seed(sender="כללית")
        seed(sender="מכבי", status="approved")
        r = client.get("/records", params={"sender": "כללית", "status": "approved"})
        assert total(r) == 1


class TestApprovedAt:
    def test_a_pending_record_has_no_approval_time_and_an_approved_one_has(self, client):
        pending, approved = seed(), seed(status="approved")
        assert client.get(f"/records/{pending}").json()["approved_at"] is None
        assert client.get(f"/records/{approved}").json()["approved_at"] is not None

    def test_editing_an_approved_record_clears_the_approval_time(self, client):
        rec = seed(status="approved")
        out = client.patch(f"/records/{rec}", json={"first_name": "דני"}).json()
        assert out["status"] == "pending" and out["approved_at"] is None


def csv_rows(response) -> list[list[str]]:
    return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))


class TestExport:
    def test_csv_download(self, client):
        seed()
        r = client.get("/export", params={"format": "csv"})
        assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
        assert r.headers["content-disposition"].startswith('attachment; filename="forms-') and r.headers["content-disposition"].endswith('.csv"')
        assert r.content.startswith(b"\xef\xbb\xbf")
        rows = csv_rows(r)
        assert len(rows) == 2 and rows[1][0] == "דנה"

    def test_excel_is_the_default_and_opens(self, client):
        seed()
        r = client.get("/export")
        assert r.headers["content-type"].endswith("spreadsheetml.sheet") and r.headers["content-disposition"].endswith('.xlsx"')
        ws = load_workbook(io.BytesIO(r.content)).active
        assert ws.max_row == 2 and ws["C2"].value == "011223344"

    def test_an_unknown_format_is_rejected(self, client):
        assert client.get("/export", params={"format": "pdf"}).status_code == 422

    def test_it_respects_the_filters(self, client):
        seed(first="רונן", status="approved")
        seed(first="דנה")
        rows = csv_rows(client.get("/export", params={"format": "csv", "status": "approved"}))
        assert [r[0] for r in rows[1:]] == ["רונן"]
        rows = csv_rows(client.get("/export", params={"format": "csv", "q": "דנה"}))
        assert [r[0] for r in rows[1:]] == ["דנה"]

    def test_nothing_matching_gives_only_the_header(self, client):
        seed()
        assert len(csv_rows(client.get("/export", params={"format": "csv", "q": "zzz"}))) == 1

    @pytest.mark.parametrize("params", [
        {}, {"status": "approved"}, {"status": "pending"}, {"sender": "מכבי"}, {"q": "דנה"},
        {"from": "2026-06-12"}, {"status": "approved", "sender": "כללית"},
    ])
    def test_the_export_and_the_list_always_agree(self, client, params):
        seed(first="דנה", sender="כללית", status="approved", received="10/06/2026")
        seed(first="רונן", sender="מכבי", received="15/06/2026")
        seed(first="דנה", sender="מכבי", status="approved", received="20/06/2026")
        listed = client.get("/records", params=params)
        exported = csv_rows(client.get("/export", params={**params, "format": "csv"}))
        assert len(exported) - 1 == total(listed) == len(listed.json())

    def test_the_export_is_not_limited_to_one_page_of_the_list(self, client):
        for _ in range(7):
            seed()
        assert len(csv_rows(client.get("/export", params={"format": "csv"}))) == 8
