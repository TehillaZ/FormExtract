import csv
import io
from datetime import UTC, date, datetime

import pytest
from openpyxl import load_workbook

from app.pipeline.extract import build_record
from app.schemas import RawExtraction, RawValue as V
from app.services.export import HEADERS, safe_text, to_csv, to_xlsx


def record(first="דנה", last="כהן", id_="001122344", central=("939376",), commitment=("730578626",),
           sender="כללית", received="18/06/2026", status="pending", approved_at=None):
    raw = RawExtraction(
        first_name=V(value=first, source="printed") if first else V(),
        last_name=V(value=last, source="printed") if last else V(),
        id_numbers=[V(value=id_, source="printed")] if id_ else [],
        central_numbers=[V(value=c, source="sticker") for c in central],
        commitment_numbers=[V(value=c, source="printed") for c in commitment],
        sender=V(value=sender, source="printed", highlighted=True) if sender else V(),
        stamp_date=V(value=received, source="stamp") if received else V(),
    )
    rec = build_record(raw, ["939376", "730578626"])
    rec.status, rec.approved_at = status, approved_at
    return rec


def csv_rows(data: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))


def sheet(data: bytes):
    return load_workbook(io.BytesIO(data)).active


class TestCsv:
    def test_starts_with_a_byte_order_mark_so_excel_shows_hebrew(self):
        assert to_csv([record()]).startswith(b"\xef\xbb\xbf")

    def test_has_a_header_row_and_one_row_per_record(self):
        rows = csv_rows(to_csv([record(), record(first="רונן", last="פלג")]))
        assert rows[0] == HEADERS
        assert [r[0] for r in rows[1:]] == ["דנה", "רונן"]

    def test_row_content(self):
        approved = datetime(2026, 6, 20, 9, 30, 45, 123, tzinfo=UTC)
        row = csv_rows(to_csv([record(status="approved", approved_at=approved)]))[1]
        assert row[:9] == ["דנה", "כהן", "001122344", "939376", "2026-06-18", "כללית", "730578626", "מאושר", "כן"]
        assert row[9] == "2026-06-20 09:30"          # UTC, to the minute
        assert len(row[10]) == 36                    # the record id

    def test_several_numbers_share_a_cell(self):
        row = csv_rows(to_csv([record(central=("939376", "939377"), commitment=("730578626", "X132767"))]))[1]
        assert row[3] == "939376; 939377" and row[6] == "730578626; X132767"

    def test_missing_values_are_empty_cells(self):
        row = csv_rows(to_csv([record(first=None, last=None, id_=None, central=(), commitment=(), sender=None, received=None)]))[1]
        assert row[:7] == ["", "", "", "", "", "", ""]

    def test_pending_has_no_approval_time(self):
        assert csv_rows(to_csv([record()]))[1][9] == ""

    def test_no_records_gives_just_the_header(self):
        assert csv_rows(to_csv([])) == [HEADERS]

    @pytest.mark.parametrize("text", ["=HYPERLINK(\"http://evil\",\"x\")", "+1+1", "-2", "@SUM(A1)"])
    def test_text_that_looks_like_a_formula_is_defused(self, text):
        assert safe_text(text) == "'" + text
        assert csv_rows(to_csv([record(first=text)]))[1][0] == "'" + text

    @pytest.mark.parametrize("text", ["\t=1", "\r=1"])
    def test_a_leading_tab_or_carriage_return_is_defused_too(self, text):
        # extraction already strips leading whitespace from names, but the export does not rely on that
        assert safe_text(text) == "'" + text

    def test_ordinary_text_is_left_alone(self):
        assert safe_text("דנה") == "דנה" and safe_text("") == ""


class TestXlsx:
    def test_is_a_right_to_left_sheet_with_a_bold_frozen_header(self):
        ws = sheet(to_xlsx([record()]))
        assert ws.sheet_view.rightToLeft is True
        assert [c.value for c in ws[1]] == HEADERS
        assert ws[1][0].font.bold and ws.freeze_panes == "A2"

    def test_ids_and_numbers_keep_their_leading_zeros_as_text(self):
        ws = sheet(to_xlsx([record(id_="001122344", central=("093376",))]))
        id_cell, central_cell = ws["C2"], ws["D2"]
        assert id_cell.value == "001122344" and id_cell.data_type == "s" and id_cell.number_format == "@"
        assert central_cell.value == "093376" and central_cell.data_type == "s"

    def test_dates_are_real_dates_shown_as_dd_mm_yyyy(self):
        approved = datetime(2026, 6, 20, 9, 30, tzinfo=UTC)
        ws = sheet(to_xlsx([record(status="approved", approved_at=approved)]))
        assert ws["E2"].value == datetime(2026, 6, 18) and ws["E2"].number_format == "DD/MM/YYYY"
        assert ws["J2"].value == datetime(2026, 6, 20, 9, 30) and ws["J2"].number_format == "DD/MM/YYYY HH:MM"

    def test_text_that_looks_like_a_formula_stays_text(self):
        ws = sheet(to_xlsx([record(first="=1+1", last="@x")]))
        assert ws["A2"].value == "=1+1" and ws["A2"].data_type == "s"     # not a formula
        assert ws["B2"].value == "@x" and ws["B2"].data_type == "s"

    def test_missing_dates_are_empty(self):
        ws = sheet(to_xlsx([record(received=None)]))
        assert ws["E2"].value is None and ws["J2"].value is None

    def test_no_records_gives_just_the_header(self):
        assert sheet(to_xlsx([])).max_row == 1

    def test_same_rows_as_the_csv(self):
        recs = [record(), record(first="רונן", last="פלג", central=("939158",))]
        ws = sheet(to_xlsx(recs))
        csv_data = csv_rows(to_csv(recs))
        for row in range(2, 4):
            assert [ws.cell(row, c).value for c in (1, 2, 3, 4, 6, 7, 8, 9)] == [csv_data[row - 1][i] for i in (0, 1, 2, 3, 5, 6, 7, 8)]
