"""Records as CSV or Excel (PLAN 9, extra 1).

Two things matter more than they look:
  * Text read from a form is untrusted. Excel runs a cell that starts with = + - @ as a formula, so in CSV such
    text is prefixed with an apostrophe, and in Excel files every text cell is stored as text, never as a formula.
  * An ID like 001122344 loses its leading zeros when Excel reads it as a number. Excel files store IDs and
    numbers as text. CSV cannot carry that, so CSV is meant for programs and the Excel file for people.
"""
from __future__ import annotations

import csv
import io
from datetime import UTC, date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from app.schemas import CommitmentRecord, Field

HEADERS = ["שם פרטי", "שם משפחה", 'ת"ז', "מספר מרכזי", "תאריך קבלה", "גורם שולח", "מספר התחייבות",
           "סטטוס", "דורש בדיקה", "תאריך אישור", "מזהה רשומה"]
TEXT_COLUMNS = {0, 1, 2, 3, 5, 6, 7, 8, 10}          # everything except the two date columns
COLUMN_WIDTHS = [14, 16, 14, 16, 14, 28, 20, 10, 12, 18, 38]
STATUS_LABEL = {"pending": "ממתין", "approved": "מאושר"}
FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")


def safe_text(text: str) -> str:
    """CSV only: keep a spreadsheet from running text that starts like a formula."""
    return "'" + text if text.startswith(FORMULA_STARTS) else text


def _joined(fields: list[Field]) -> str:
    return "; ".join(f.value for f in fields if f.value)


def _received(rec: CommitmentRecord) -> date | None:
    return date.fromisoformat(rec.received_date.value) if rec.received_date.value else None


def _approved(rec: CommitmentRecord) -> datetime | None:
    """When it was approved, as a UTC time without a zone (spreadsheets have no time zones)."""
    if rec.approved_at is None:
        return None
    moment = rec.approved_at if rec.approved_at.tzinfo else rec.approved_at.replace(tzinfo=UTC)
    return moment.astimezone(UTC).replace(tzinfo=None, microsecond=0)


def _values(rec: CommitmentRecord) -> list:
    """One row in the order of HEADERS. Dates stay dates; the caller decides how to write them."""
    return [
        rec.first_name.value or "", rec.last_name.value or "", rec.id_number.value or "",
        _joined(rec.central_numbers), _received(rec), rec.sender.value or "", _joined(rec.commitment_numbers),
        STATUS_LABEL.get(rec.status, rec.status), "כן" if rec.needs_review() else "לא", _approved(rec), str(rec.id),
    ]


def to_csv(records: list[CommitmentRecord]) -> bytes:
    """UTF-8 with a byte-order mark, which is what makes Excel show Hebrew correctly. Dates are ISO (unambiguous)."""
    out = io.StringIO(newline="")
    writer = csv.writer(out)
    writer.writerow(HEADERS)
    for rec in records:
        row = []
        for i, value in enumerate(_values(rec)):
            if isinstance(value, datetime):
                row.append(value.strftime("%Y-%m-%d %H:%M"))
            elif isinstance(value, date):
                row.append(value.isoformat())
            else:
                row.append(safe_text(value) if i in TEXT_COLUMNS else value)
        writer.writerow(row)
    return ("﻿" + out.getvalue()).encode("utf-8")


def to_xlsx(records: list[CommitmentRecord]) -> bytes:
    """A right-to-left sheet. Text columns are real text (leading zeros kept, nothing runs as a formula);
    dates are real dates shown as dd/mm/yyyy."""
    wb = Workbook()
    ws = wb.active
    ws.title = "רשומות"
    ws.sheet_view.rightToLeft = True
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for rec in records:
        ws.append(_values(rec))
        for i, cell in enumerate(ws[ws.max_row]):
            if i in TEXT_COLUMNS:
                cell.number_format = "@"
                cell.data_type = "s"          # even "=1+1" stays text
            elif isinstance(cell.value, datetime):
                cell.number_format = "DD/MM/YYYY HH:MM"
            elif isinstance(cell.value, date):
                cell.number_format = "DD/MM/YYYY"
    for column, width in zip("ABCDEFGHIJK", COLUMN_WIDTHS):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
